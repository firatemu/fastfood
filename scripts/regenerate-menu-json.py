import openpyxl, re, unicodedata, json

XLSX = "/home/azem/projects/fastfood/Burger_Home_Menu.xlsx"
OUT = "/home/azem/projects/fastfood/apps/api/prisma/seed-data/menu-xlsx.json"
NUM = re.compile(r"^(\d+)\s*-\s*")
TR = [("ı","i"),("İ","i"),("ş","s"),("Ş","s"),("ğ","g"),("Ğ","g"),
      ("ü","u"),("Ü","u"),("ö","o"),("Ö","o"),("ç","c"),("Ç","c")]

def slugify(s):
    s = s.replace("'", "").replace("’", "")          # 2'LI -> 2li (seed convention)
    for a, b in TR: s = s.replace(a, b)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-zA-Z0-9]+", "-", s.lower()).strip("-")

def clean_desc(s):
    if not s: return None
    s = str(s)
    s = s.replace("fleto", "fileto").replace("coleslow", "coleslaw")
    s = s.replace("%21", "%").replace("Fleto", "Fileto")
    s = re.sub(r"\s*,\s*", ", ", s)
    s = re.sub(r"\s{2,}", " ", s).strip(" ,")
    return s or None

def minor(v):
    return int(round(float(v.replace(",", ".")) * 100))

# ── price cell -> [(label, minor)] ───────────────────────────────────────────
def parse_prices(cell):
    """'130 GR.: 464.95 ₺ / 100 GR.: 429.95 ₺' -> [('130 gr',46495),('100 gr',42995)]"""
    if not cell: return []
    out = []
    for part in str(cell).split("/"):
        m = re.search(r"([0-9]+(?:[.,][0-9]+)?)\s*₺", part)
        if not m: continue
        label = part[: m.start()]
        g = re.match(r"\s*(\d+)\s*(?:GR\.?|gr\.?)?\s*:?\s*$", label)
        if g:
            lbl = f"{int(g.group(1))} gr"
        else:
            lbl = part.strip(" .:") or "Standart"
        out.append((lbl, minor(m.group(1))))
    return out

# ── canonical extras: slugified key -> (group, display name, price minor) ────
EXTRAS = {
    "100-gr-dana-eti":     ("İlave Köfte",    "100 gr Dana Köfte",   11995),
    "260gr-dana-eti":      ("İlave Köfte",    "260 gr Dana Köfte",   12695),
    "cheddar-peynir":      ("Ekstra Peynir",  "Cheddar Peynir",      1945),
    "mozzarella-peynir":   ("Ekstra Peynir",  "Mozzarella Peynir",   1945),
    "hellim-peynir":       ("Ekstra Peynir",  "Hellim Peynir",       1945),
    "bacon":               ("Ekstra Malzeme", "Bacon",                3295),
    "soslu-mantar":        ("Ekstra Malzeme", "Soslu Mantar",         2045),
    "coleslaw-salata":     ("Ekstra Malzeme", "Coleslaw Salata",      2145),
    "coleslaw-salatasi":   ("Ekstra Malzeme", "Coleslaw Salata",      2145),
    "kitir-sogan":         ("Ekstra Malzeme", "Kıtır Soğan",          1845),
    "kibrit-patates":      ("Ekstra Malzeme", "Kibrit Patates",       1695),
    "yumurta":             ("Ekstra Malzeme", "Yumurta",              1645),
    "cheddar-sos":         ("Soslar",         "Cheddar Sos",          4995),
    "truf-sos":            ("Soslar",         "Trüf Sos",             2295),
    "crispy-sos":          ("Soslar",         "Crispy Sos",           2045),
    "buyuk-boy-patates":   ("Patates Boyutu", "Büyük Boy Patates",    2345),
}
def canon_extra(raw):
    s = re.sub(r"\(.*?\)", " ", str(raw))                   # "(burger içinde)"
    s = re.sub(r"burger\s*içi\w*", " ", s, flags=re.I)      # "burger içi"
    return EXTRAS.get(slugify(s).strip("-"))

# ── read workbook ────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["Tüm Menü (Özet)"]
ws2 = wb["Ekstra Detay Listesi"]

raw_extras = {}
for r in range(2, ws2.max_row + 1):
    p = ws2.cell(row=r, column=2).value
    e = ws2.cell(row=r, column=5).value
    if p is None: continue
    raw_extras.setdefault(str(p).strip(), set()).add(str(e).strip() if e else "")

cat, categories, products, order = None, [], [], {}
for r in range(2, ws.max_row + 1):
    a, b, c, d = (ws.cell(row=r, column=i).value for i in range(1, 5))
    if b is None:
        cat = str(a).strip()
        categories.append({"slug": slugify(cat), "name": cat, "displayOrder": len(categories) + 1})
        order[cat] = 0
        continue
    excel_name = str(b).strip()
    m = NUM.match(excel_name)
    clean = NUM.sub("", excel_name).strip()
    slug = (f"{m.group(1)}-" if m else "") + slugify(clean)
    order[cat] += 1
    prices = parse_prices(d)
    prices.sort(key=lambda t: t[1])
    groups = {}
    unknown = set()
    for raw in raw_extras.get(excel_name, set()):
        if not raw or raw.strip().lower() == "yok": continue
        hit = canon_extra(raw)
        if hit: groups.setdefault(hit[0], []).append({"name": hit[1], "priceMinor": hit[2]})
        else: unknown.add(raw)
    if unknown: print(f"  ⚠ eşleşmeyen ekstra [{excel_name}]: {sorted(unknown)}")
    products.append({
        "excelName": excel_name, "name": clean, "slug": slug,
        "categorySlug": slugify(cat), "displayOrder": order[cat],
        "description": clean_desc(c),
        "prices": [{"label": lbl, "priceMinor": val} for lbl, val in prices],
        "modifierGroups": {k: sorted(v, key=lambda x: x["name"]) for k, v in sorted(groups.items())},
    })

payload = {
    "_source": "Burger_Home_Menu.xlsx",
    "_generatedBy": "apps/api/prisma/seed-menu-xlsx.ts tarafından okunur; scripts/regenerate-menu-json.py ile üretilir",
    "groups": {
        "İlave Köfte":    {"selectionMin": 0, "selectionMax": 1},
        "Ekstra Peynir":  {"selectionMin": 0, "selectionMax": 3},
        "Ekstra Malzeme": {"selectionMin": 0, "selectionMax": 6},
        "Soslar":         {"selectionMin": 0, "selectionMax": 1},
        "Patates Boyutu": {"selectionMin": 1, "selectionMax": 1,
                           "required": [{"name": "Standart", "priceMinor": 0}]},
    },
    "categories": categories,
    "products": products,
}
json.dump(payload, open(OUT, "w"), ensure_ascii=False, indent=1)

print(f"\nkategori: {len(categories)}  ürün: {len(products)}")
print(f"-> {OUT}")
print("\nvaryantlı ürün örneği:")
for p in products[:2] + [x for x in products if x["categorySlug"] == "tavuk-burgerler"][:1]:
    print(f"  {p['name']:<32} base={p['prices'][0]}  gruplar={list(p['modifierGroups'])}")